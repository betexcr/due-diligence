"""Document processing: turn stored bytes into page/sheet-addressed text so every
extracted value can be traced back to where it came from."""

from __future__ import annotations

from pathlib import Path

from .contracts import DocumentContent, PageText
from .errors import PermanentError

PDF_TYPES = {"application/pdf"}
XLSX_TYPES = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
TEXT_TYPES = {"text/plain", "text/markdown", "text/csv"}
ALLOWED_TYPES = PDF_TYPES | XLSX_TYPES | TEXT_TYPES

EXTENSION_TYPES = {
    ".pdf": "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
}

# Below this many characters per page we assume a scanned PDF that needs OCR.
MIN_TEXT_CHARS_PER_PAGE = 20


def sniff_content_type(file_name: str, declared: str | None, head: bytes) -> str:
    """Trust the bytes over the browser's declared type."""
    if head.startswith(b"%PDF-"):
        return "application/pdf"
    if head.startswith(b"PK\x03\x04") and file_name.lower().endswith(".xlsx"):
        return EXTENSION_TYPES[".xlsx"]
    ext = Path(file_name).suffix.lower()
    if ext in (".txt", ".md", ".csv"):
        return EXTENSION_TYPES[ext]
    return declared or "application/octet-stream"


def extract_content(path: Path, content_type: str) -> DocumentContent:
    if content_type in PDF_TYPES:
        return _pdf(path)
    if content_type in XLSX_TYPES:
        return _xlsx(path)
    if content_type in TEXT_TYPES:
        return DocumentContent(kind="text", pages=[PageText(page=1, text=path.read_text("utf-8"))])
    raise PermanentError(f"Unsupported content type: {content_type}")


def _pdf(path: Path) -> DocumentContent:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(path)
        pages = [PageText(page=i + 1, text=(p.extract_text() or "").strip())
                 for i, p in enumerate(reader.pages)]
    except (PdfReadError, ValueError) as exc:
        raise PermanentError(f"Corrupt or unreadable PDF: {exc}") from exc
    if not pages:
        raise PermanentError("PDF has no pages")
    if sum(len(p.text) for p in pages) < MIN_TEXT_CHARS_PER_PAGE * len(pages):
        # Production: route to OCR (Textract / Azure Document Intelligence / Tesseract)
        # or send page images to Claude directly. Out of scope for the prototype.
        raise PermanentError("PDF has no embedded text layer (scanned?). OCR required.")
    return DocumentContent(kind="pdf", pages=pages)


def _xlsx(path: Path) -> DocumentContent:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        raise PermanentError(f"Corrupt or unreadable spreadsheet: {exc}") from exc
    pages = []
    for ws in wb.worksheets:
        lines = []
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}={c.value}" for c in row
                     if c.value is not None and hasattr(c, "coordinate")]
            if cells:
                lines.append(" | ".join(cells))
        if lines:
            pages.append(PageText(sheet=ws.title, text="\n".join(lines)))
    wb.close()
    if not pages:
        raise PermanentError("Spreadsheet is empty")
    return DocumentContent(kind="xlsx", pages=pages)
